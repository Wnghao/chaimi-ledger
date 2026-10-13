"""账单导入：支付宝/微信/京东 + 本应用导出（CSV/xlsx）解析 + 三段式分类 + 退款匹配 + 去重（DESIGN 七/八章）。

链路：parse_bill(文件) → FlowPreview[]（不写库）→ apply_import(确认) 。
金额全程整数分。
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import sqlite3
from datetime import datetime

import db as dbm

# ---------- 基础工具 ----------

_JD_REFUND_DONE_STATUSES = frozenset({
    "退款成功", "已退款", "退款完成", "全额退款", "部分退款",
    "退款已到账", "退款到账", "成功退款", "退已成功",
})

#   匹配含"退款"关键词的括号，支持以下形式：
#   "(已退款X)" / "(已全额退款)" / "(部分退款X)" / "(退款成功X)" / "(退款完成X)" / "(退款到账X)"
#   括号内数字作为退款金额；无数字（如已全额退款）则退款金额=原金额
#   仅匹配半角会导致全角行直接 float() 抛错
_JD_AMOUNT_REFUND_RE = re.compile(r"[（(]([^（）()]*退款[^（）()]*)[)）]")

#   用于"已退积分"/"红包"/"入账"/"补偿"/"不计收支"等非退款特殊词语
#   匹配后标 need_confirm=1 走待确认，不入库退款流水
#   兼容半角/全角括号（理由同上）
_JD_AMOUNT_PAREN_RE = re.compile(r"[（(]([^（）()]*)[)）]")

def normalize(field) -> str:
    """trim（含京东单元格尾部 \t）。"""
    return str(field or "").strip().replace("\t", "")

def to_cents(s) -> int:
    s = normalize(s).replace("¥", "").replace("￥", "").replace(",", "")
    return round(float(s or "0") * 100)

def parse_dt(s: str) -> str:
    s = normalize(s).replace("/", "-")
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            continue
    return s

def _row_hash(fields: list[str]) -> str:
    return hashlib.md5("|".join(normalize(f) for f in fields).encode("utf-8")).hexdigest()

# 解析期行数熔断上限（防高压缩比 xlsx / 超长 CSV 的解析期内存放大）；
# server.import_parse 的 len(flows) 校验引用同一常量作权威上限
MAX_PARSE_ROWS = 50_000

def _capped(rows):
    """行迭代熔断：超限抛 ValueError（import_parse 统一转 400）。
    余量 200 行给各平台表头前的导出信息块。"""
    for i, row in enumerate(rows):
        if i > MAX_PARSE_ROWS + 200:
            raise ValueError("行数超过上限 5 万行")
        yield row

def _hash_fields(get, keys: list[str], amt_key: str, amt_cents: int | None = None) -> list:
    """构造 hash 字段列表（各平台解析器共用）：金额列(amt_key)规范化为整数分字符串
    （amt_cents 已给出则直接用，如京东拆括号后的原金额），其余列原样。"""
    return [(str(amt_cents if amt_cents is not None else to_cents(get(k)))
             if k == amt_key else get(k)) for k in keys]

def _find_header(rows: list[list], key: str) -> int:
    for i, r in enumerate(rows):
        if r and normalize(r[0]) == key:
            return i
    return -1

def _mk_getter(idx: dict, row: list):
    """按表头索引取单元格；行短时回退空串（各平台解析器共用）。"""
    return lambda k: row[idx[k]] if idx[k] < len(row) else ""

def _note_fallback(idx: dict, header: list) -> dict:
    """备注列回退（支付宝/京东共用）：表头无「备注」时借最后一列（备注为可选导出项）。"""
    idx["备注"] = header.index("备注") if "备注" in header else len(header) - 1
    return idx

def _ttype(shouzhi: str, is_refund: bool) -> str:
    """收/支列 → 流水类型（支付宝/微信共用）：退款强制 transfer，其余按方向，未标注为 transfer。"""
    if is_refund:
        return "transfer"
    if shouzhi == "支出":
        return "expense"
    if shouzhi == "收入":
        return "income"
    return "transfer"

# ---------- 平台解析器 → RawFlow ----------

class RawFlow(dict):
    """平台无关的原始流水。"""

def _join_note(*parts) -> str:
    """备注拼接：商户/对方 + 备注列，
    非空部分空格连接。"""
    return " ".join(p for p in (normalize(x) for x in parts) if p)

def _is_placeholder_name(name: str) -> bool:
    """占位名判定：微信账单「商品」列无内容时常写作 "/"（红包/转账行普遍如此），这类占位名
    不参与 ② 的名称匹配：否则两条同样写作 "/" 的行会因「金额相等 + 前缀相等」被判为高置信，
    把退款关联到无关的红包/转账行。"""
    return not name.strip("/-—\\ \u3000")

WX_GENERIC_NAMES = frozenset({
    "微信红包", "微信红包（单发）", "微信红包（群红包）", "群红包", "红包",
    "微信转账", "转账", "微信红包-退款", "转账-退款",
})

def is_meaningless_name(name: str) -> bool:
    """名称是否无意义：空 / 占位符（"/"、"-"）/ 微信红包·转账这类泛化值。
    用于「对方优先 + AI 兜底」的名称回填判定。"""
    n = normalize(name)
    return _is_placeholder_name(n) or n in WX_GENERIC_NAMES

def wx_name(product: str, merchant: str) -> str:
    """微信名称回填：仅当「商品」列无意义时，用「交易对方」顶替；
    对方也是泛化值（微信红包/转账/空）时留空，由 ai.resolve_names_batch 读备注等兜底提取。"""
    n = normalize(product)
    if not is_meaningless_name(n):
        return n
    m = normalize(merchant)
    return "" if is_meaningless_name(m) else m

# 账单主人识别：三平台导出文件的「导出信息」块里各有账号标识行
# （微信：微信昵称：xxx / 京东：京东账号名：xxx / 支付宝：姓名：xxx）。
# 识别结果作为整份账单每笔流水的流水归属（tags.attribution），一家人各自的账单
# 混在同一个账本时即可按人区分。姓名放最后兜底（支付宝块里同时有「支付宝账户」
# 邮箱行，不在捕获键里，不会误抓）。
_OWNER_RE = re.compile(r"(微信昵称|京东账号名|姓名)\s*[:：]\s*([^\s，,]+)")

def _detect_owner(rows, h) -> str:
    """在表头行之前的导出信息里找账号主人；找不到返回空串（归属留空）。"""
    for r in rows[:h]:
        for cell in r:
            m = _OWNER_RE.search(str(cell or ""))
            if m:
                return m.group(2).strip()
    return ""

def parse_alipay(text: str) -> list[RawFlow]:
    rows = list(_capped(csv.reader(io.StringIO(text))))
    h = _find_header(rows, "交易时间")
    if h < 0:
        raise ValueError("未找到支付宝表头（交易时间）")
    idx = _note_fallback({k: rows[h].index(k) for k in
                          ["交易时间", "交易分类", "交易对方", "商品说明", "收/支", "金额",
                           "收/付款方式", "交易状态", "交易订单号", "商家订单号"]}, rows[h])
    owner = _detect_owner(rows, h)      # 支付宝：姓名
    out = []
    for r in rows[h + 1:]:
        if not r or not normalize(r[0]):
            continue
        g = _mk_getter(idx, r)
        # 不过滤交易状态：「交易关闭」实测多为已全额退款的原消费（退款行独立存在），
        # 保留才能让退款匹配/统计抵减闭环（净额=0）；0 元行（如医保支付）跳过。
        cat, shouzhi = normalize(g("交易分类")), normalize(g("收/支"))
        amount = to_cents(g("金额"))
        is_refund = cat == "退款"
        if amount == 0:
            out.append(RawFlow(skip="0 元退款行"))
            continue
        ttype = _ttype(shouzhi, is_refund)
        out.append(RawFlow(
            platform="alipay", date=parse_dt(g("交易时间")), source_category=cat,
            merchant=normalize(g("交易对方")), name=normalize(g("商品说明")),
            amount=amount, type=ttype, is_refund=is_refund, attribution=owner,
            pay_account=normalize(g("收/付款方式")),
            order_no=normalize(g("交易订单号")),
            note=_join_note(g("交易对方"), g("备注")),
            hash_fields=_hash_fields(g, ["交易时间", "交易分类", "交易对方", "商品说明",
                                         "金额", "商家订单号", "备注"], "金额")))
    return out

# 但本质是自有账户间搬钱（还信用卡、买理财/余额宝、给钱包充值），算收支会与刷卡消费
# 重复计算 → 按 TRANSFER_SUBCLASS 原意归 transfer 不计收支。
# 从 TRANSFER_SUBCLASS 派生，排除「账户转账」族（个人间转账按方向计收支）与退款（单独处理）
_WX_TRANSFER_ONLY = frozenset(dbm.TRANSFER_SUBCLASS) - {
    "转账", "转账红包", "微信红包", "微信红包（单发）", "退款",
}

def parse_wechat(wb) -> list[RawFlow]:
    ws = wb.active
    rows = [list(r) for r in _capped(ws.iter_rows(values_only=True))]
    h = _find_header(rows, "交易时间")
    if h < 0:
        raise ValueError("未找到微信表头（交易时间）")
    idx = {k: rows[h].index(k) for k in
           ["交易时间", "交易类型", "交易对方", "商品", "收/支", "金额(元)",
            "支付方式", "当前状态", "交易单号", "商户单号"]}
    # 备注列单独取索引、不进 idx：微信备注不参与 row_hash（有无备注列的两次导出仍判同一行）
    note_col = rows[h].index("备注") if "备注" in rows[h] else -1
    owner = _detect_owner(rows, h)      # 微信：微信昵称
    out = []
    for r in rows[h + 1:]:
        if not r or not normalize(r[0]):
            continue
        g = _mk_getter(idx, r)
        g_note = lambda: r[note_col] if 0 <= note_col < len(r) else ""
        status = normalize(g("当前状态"))
        if status in ("交易关闭", "已删除"):
            continue
        ttype_raw, shouzhi = normalize(g("交易类型")), normalize(g("收/支"))
        amount = to_cents(g("金额(元)"))
        # "已全额退款"（实测红包：原消费行"微信红包（单发）/支出/已全额退款"），
        # 若不带方向判断会把原消费误判成退款行 → 强制 transfer → 既不计收支、又无法被
        # 退款行匹配（_match_refunds 的候选只收 type='expense'）→ 关联下拉里找不到它。
        # 退款一定表现为「钱回到我账户」（收入），故状态判定仅在非支出方向生效。
        is_refund = ttype_raw.endswith("-退款") or (
            status in ("已全额退款", "对方已退还") and shouzhi != "支出")
        base_type = ttype_raw[:-3] if ttype_raw.endswith("-退款") else ttype_raw
        #   ① 个人间转账/红包按方向计收支：我转给他人=支出、他人转我=收入
        #   ② 账户间搬运（信用卡还款/理财申赎/充值）不计收支 → transfer
        #   ③「收/支」为空的中性交易（零钱提现、零钱通转入转出等）也留作 transfer
        if is_refund:
            ttype = "transfer"  # 微信退款标"收入"，强制 transfer（7.4）
        elif base_type in _WX_TRANSFER_ONLY:
            ttype = "transfer"  # 账户间搬运不计收支（信用卡还款等）
        else:
            ttype = _ttype(shouzhi, False)
        out.append(RawFlow(
            platform="wechat", date=parse_dt(g("交易时间")), source_category=base_type,
            merchant=normalize(g("交易对方")), name=wx_name(g("商品"), g("交易对方")),
            amount=amount, type=ttype, is_refund=is_refund, attribution=owner,
            pay_account=normalize(g("支付方式")),
            order_no=normalize(g("交易单号")),
            note=_join_note(g("交易对方"), g_note()), wechat_type=base_type,
            hash_fields=_hash_fields(g, ["交易时间", "交易类型", "交易对方", "商品",
                                         "金额(元)", "交易单号"], "金额(元)")))
    return out

def parse_jd(text: str) -> list[RawFlow]:
    rows = list(_capped(csv.reader(io.StringIO(text))))
    h = _find_header(rows, "交易时间")
    if h < 0:
        raise ValueError("未找到京东表头（交易时间）")
    idx = _note_fallback({k: rows[h].index(k) for k in
                          ["交易时间", "商户名称", "交易说明", "金额", "收/支", "交易状态",
                           "交易分类", "交易订单号", "收/付款方式"]}, rows[h])
    out = []
    # 同订单号+同时间但不同金额的两笔交易被错误去重。金额列写法差异
    # （"100"/"100.00"/"100(已退款30)" 拆括号后等）由 to_cents 在调用方规范化后传入，
    # 保证 hash 字段中的金额是整数 cents 字符串而非原始单元格）
    jd_hash_keys = ["交易时间", "交易分类", "商户名称", "交易说明", "金额", "交易订单号", "备注"]
    owner = _detect_owner(rows, h)      # 京东：京东账号名
    def _jd_hash(g, extra="", amt_cents=None):
        # amt_cents：整数 cents（已规范化）；未传则由 _hash_fields 从 g("金额") 原始值计算
        return _hash_fields(g, jd_hash_keys, "金额", amt_cents) + ([extra] if extra else [])
    for r in rows[h + 1:]:
        if not r or not normalize(r[0]):
            continue
        g = _mk_getter(idx, r)
        shouzhi, status = normalize(g("收/支")), normalize(g("交易状态"))
        raw_amount = str(g("金额") or "").strip().replace("\t", "")
        is_refund = status in _JD_REFUND_DONE_STATUSES or (
            "退款" in status and "处理中" not in status and "退款中" != status
        )
        cats = normalize(g("交易分类")).split()  # 京东空格分隔多分类
        base = dict(
            platform="jd", date=parse_dt(g("交易时间")),
            source_category=cats[0] if cats else "",
            extra_categories=cats[1:], merchant=normalize(g("商户名称")),
            pay_account=normalize(g("收/付款方式")),
            order_no=normalize(g("交易订单号")),
            note=_join_note(g("商户名称"), g("备注")),
            attribution=owner,
        )
        name = normalize(g("交易说明"))
        #   ①含"退款"关键词的括号 → 拆为原消费 + 退款两条流水（_match_refunds 通过 order_no 自动高置信关联）
        #   ②含其他括号内容（已退积分/红包/入账等非退款特殊词）→ 单条流水 + need_confirm=1
        #   ③纯数字 → 走原 ttype 判断
        amount_m = _JD_AMOUNT_REFUND_RE.search(raw_amount)
        if amount_m:
            inner = amount_m.group(1)
            origin_amt_str = _JD_AMOUNT_REFUND_RE.sub("", raw_amount)
            # 提取括号内最后一个数字作为退款金额（兼容"已退款50.08"/"退款成功30"等）
            # 无数字（如"已全额退款"）则退款金额=原金额 → 净额自动全额抵销
            nums = re.findall(r"\d+(?:\.\d+)?", inner)
            refund_amt_str = nums[-1] if nums else origin_amt_str
            origin_cents = to_cents(origin_amt_str)
            refund_cents = to_cents(refund_amt_str)
            base_hash = _jd_hash(g, amt_cents=origin_cents)
            # 原消费行：括号外的金额，强制 expense
            out.append(RawFlow(**base, name=name, amount=origin_cents,
                               type="expense", is_refund=False, hash_fields=base_hash))
            # 退款行：括号内金额，type=transfer；hash 末位加 jd_refund:<cents> 区分避免去重误判
            # 退款行 hash 金额字段复用 origin_cents，与原消费行 hash 保持一致，唯一区别靠
            # jd_refund extra 区分，避免原始单元格写法（"100(已退款30)" vs "100.00(已退款30.00)"）
            # 导致同一笔订单的两条流水 hash 不一致）
            out.append(RawFlow(**base, name=name, amount=refund_cents,
                               type="transfer", is_refund=True,
                               hash_fields=base_hash + [f"jd_refund:{refund_cents}"]))
            continue
        # 非退款的括号内容（已退积分/红包/入账/补偿/不计收支等）→ 单条流水 + 待确认
        paren_m = _JD_AMOUNT_PAREN_RE.search(raw_amount)
        if paren_m:
            origin_amt_str = _JD_AMOUNT_PAREN_RE.sub("", raw_amount).strip()
            origin_cents = to_cents(origin_amt_str) if origin_amt_str else 0
            # 按 收/支 列判断 ttype；京东"不计收支"列值走 transfer
            ttype = "expense" if shouzhi == "支出" else "transfer"
            out.append(RawFlow(**base, name=name, amount=origin_cents,
                               type=ttype, is_refund=False, need_confirm=True,
                               hash_fields=_jd_hash(g, amt_cents=origin_cents)))
            continue
        # is_refund 强制 transfer；否则按"收/支"列判断（"支出"→expense，其余含"不计收支"→transfer）
        ttype = "expense" if not is_refund and shouzhi == "支出" else "transfer"
        amount_cents = to_cents(raw_amount)
        out.append(RawFlow(**base, name=name, amount=amount_cents,
                           type=ttype, is_refund=is_refund,
                           hash_fields=_jd_hash(g, amt_cents=amount_cents)))
    return out

def parse_app_export(text: str) -> list[RawFlow]:
    """本应用导出格式。导出回环三件套：
    收款账户列（转账对方，旧版两账户 & 拼接自动拆）、订单号/来源列原样还原、
    流水ID/退款关联列——退款对按 ID 批内直连，无需重新匹配。"""
    rows = list(_capped(csv.reader(io.StringIO(text))))
    h = _find_header(rows, "交易时间")
    if h < 0:
        h = 0
    hdr = rows[h]
    out = []
    for r in rows[h + 1:]:
        if not r or not normalize(r[0]):
            continue
        m = dict(zip(hdr, r))
        ltype = normalize(m.get("流水类型", ""))
        ttype = {"收入": "income", "支出": "expense"}.get(ltype, "transfer")
        pay = normalize(m.get("支付账户") or m.get("支付方式/收款方式", ""))
        to = normalize(m.get("收款账户", ""))
        if not to and " & " in pay:
            pay, to = [x.strip() for x in pay.split(" & ", 1)]
        out.append(RawFlow(
            platform="export", date=parse_dt(m.get("交易时间", "")),
            source_category=normalize(m.get("支出类型/收入类型", "")),
            merchant="", name=normalize(m.get("交易名称", "")),
            amount=to_cents(m.get("交易金额", "0")), type=ttype, is_refund=False,
            pay_account=pay, to_account_name=to,
            order_no=normalize(m.get("订单号", "")),
            origin_source=normalize(m.get("来源", "")),
            export_id=normalize(m.get("流水ID", "")),
            export_ref=normalize(m.get("退款关联", "")),
            note=normalize(m.get("备注", "")),
            attribution=normalize(m.get("流水归属", "")),
            hash_fields=_hash_fields(lambda k: m.get(k, ""),
                                     ["交易时间", "交易名称", "交易金额",
                                      "支出类型/收入类型", "备注"], "交易金额")))
    # 第二遍：退款关联 → 本批负序号（apply_import 的批内回填协议：-(目标序号+1)）
    by_id = {f.get("export_id"): i for i, f in enumerate(out) if f.get("export_id")}
    for f in out:
        ref = f.get("export_ref")
        if ref and ref in by_id:
            f["export_refund_idx"] = -(by_id[ref] + 1)
    return out

def parse_bill(data: bytes, filename: str) -> list[RawFlow]:
    """按文件特征识别平台并解析。"""
    fn = filename.lower()
    if fn.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        # 按首个 sheet 表头识别——微信独有「交易类型」，本应用导出含「流水类型/交易名称」；
        ws_head = " ".join(str(v) for row in wb.active.iter_rows(min_row=1, max_row=12, values_only=True)
                           for v in row if v is not None)
        if "交易类型" in ws_head:
            return parse_wechat(wb)
        if "流水类型" in ws_head or "交易名称" in ws_head:
            import csv as _csv
            buf = io.StringIO()
            w = _csv.writer(buf)
            for row in _capped(wb.active.iter_rows(values_only=True)):
                w.writerow(["" if v is None else v for v in row])
            return parse_app_export(buf.getvalue())
        raise ValueError("无法识别 xlsx 表头（支持：微信账单 / 本应用导出）")
    # CSV：京东/现金簿 UTF-8(BOM)，支付宝 GBK
    for enc in ("utf-8-sig", "gbk"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        # 个别导出含坏字节（如表情符号被写成非法 GBK 序列，严格解码必失败）：
        # 两种编码都按 errors=replace 降级试一遍，取替换字符 U+FFFD 最少的那种
        text, _ = min(((data.decode(e, errors="replace"), e) for e in ("utf-8-sig", "gbk")),
                      key=lambda p: p[0].count("\ufffd"))
    # 平台识别（列名互斥，按独有列优先）：京东=商户名称、支付宝=对方账号、本应用导出=流水类型
    if "商户名称" in text and "交易说明" in text:
        return parse_jd(text)
    if "对方账号" in text or ("支付宝" in text[:800] and "交易订单号" in text):
        return parse_alipay(text)
    if "流水类型" in text or "交易名称" in text:
        return parse_app_export(text)
    raise ValueError("无法识别账单平台（支持：支付宝/微信/京东 CSV / 本应用导出 CSV、xlsx）")

# ---------- 分类解析（三段式 + transfer 子类；二级分类已取消） ----------

# 一级分类关键词兜底（微信账单无分类列，靠商品/商户名识别；顺序=优先级；
# 只收高精度词——能跨类误伤的泛化词（如「店」「平台」）一律不收）
TOP_KEYWORD_LEXICON = [
    ("工资薪酬", ["工资", "薪资", "奖金", "报酬", "薪酬", "劳务"]),
    # 红包/礼金合并进「人情往来」（小额通用红包切「红包礼金」，见 dbm.is_small_generic_packet）
    ("人情往来", ["压岁", "红包", "礼金", "随礼", "人情往来",
                  "份子", "人情", "花圈", "奠仪", "伴手礼"]),
    ("美容美发", ["发艺", "理发", "美发", "美容", "美甲", "护肤"]),
    ("医疗健康", ["药房", "药店", "诊所", "医院", "牙科", "体检", "医药", "挂号"]),
    ("加油充电", ["加油", "加油站", "充电桩", "中石化", "中国石化", "电子钱包", "石油"]),
    ("爱车养车", ["停车", "ETC", "etc", "洗车", "捷停车", "速停车", "云停", "JSPAY",
                  "4S店", "维保", "保养", "汽修", "轮胎"]),
    ("交通出行", ["滴滴", "打车", "快车", "顺风车", "出租车", "网约车", "12306", "中铁网络", "铁路",
                  "高铁", "客运站", "汽车客运", "地铁", "公交", "机票", "航空", "机场", "共享单车"]),
    ("充值缴费", ["话费", "宽带", "水费", "电费", "燃气", "物业费", "有线电视", "阿里云", "卡充值",
                  "采暖", "取暖", "暖气"]),
    ("餐饮美食", ["食堂", "餐厅", "餐饮", "吃饭", "早餐", "早饭", "午饭", "午餐", "晚饭", "晚餐", "烧烤", "小吃", "奶茶", "咖啡", "早点", "外卖", "面馆", "火锅", "市集", "农贸市场", "饭馆", "快餐", "小吃部",
                  "饮品", "蛋糕", "烘焙", "甜品", "夜宵"]),   # 餐时词：AI 记一笔高频（早餐5元/打车15 这类），导入侧同样受益
    ("母婴亲子", ["奶粉", "尿裤", "尿不湿", "婴儿", "母婴"]),
    ("日用百货", ["超市", "百货", "连锁", "便利", "生鲜", "五金", "家具", "家居", "建材", "纸巾",
                  "卫生巾", "护垫", "湿巾"]),
    ("服饰装扮", ["男装", "女装", "童装", "服装", "服饰", "羽绒服", "运动鞋", "皮鞋", "鞋店"]),
    ("文化休闲", ["影院", "电影", "游戏", "图书", "KTV", "健身", "景区", "门票", "旅游"]),
    ("教育培训", ["培训", "网课", "教育", "学费", "课程"]),
    ("数码电器", ["数码", "电器", "手机", "电脑", "家电", "耳机", "充电宝"]),
]

_PLATE_RE = re.compile(
    "[京津沪渝冀豫云辽黑湘皖鲁新苏浙赣鄂桂甘晋蒙陕吉闽贵粤青藏川宁琼使领]"
    "[A-HJ-NP-Z][A-HJ-NP-Z0-9]{4,6}")

def _renqing_split(text: str, amount: int | None) -> str:
    """人情类关键词命中后的两级归类：
    小额通用红包 → 「红包礼金」；其余（大额或带明确人情含义）→ 「人情往来」。"""
    return "红包礼金" if dbm.is_small_generic_packet(text, amount) else "人情往来"

# 拼手气/群红包等确定字样的判定收口在 dbm.renqing_need_confirm（与词表常量同居 db.py）

# 加油词表单列：优先于车牌正则（带车牌的加油记录归加油充电而非爱车养车）
_FUEL_WORDS = dict(TOP_KEYWORD_LEXICON)["加油充电"]

# 产品级例外（优先级最高）：子串撞车的产品名直接定向，防词典误伤
_KW_EXCEPTIONS = [("航空杯", "日用百货")]

def _top_by_keyword(text: str, amount: int | None = None) -> str | None:
    # 优先级：产品例外 > 加油充电 > 车牌正则 > 词典顺序
    for word, cat in _KW_EXCEPTIONS:
        if word in text:
            return cat
    if any(w in text for w in _FUEL_WORDS):
        return "加油充电"
    if _PLATE_RE.search(text):
        return "爱车养车"
    for cat, words in TOP_KEYWORD_LEXICON:
        if any(w in text for w in words):
            # 人情类词表统一按金额/字样分到「红包礼金」或「人情往来」
            return _renqing_split(text, amount) if cat == "人情往来" else cat
    return None

def _resolve_by_mapping(conn: sqlite3.Connection, ledger_id: int, flow_type: str, source: str):
    """映射解析（v6 外键版）：source → 分类行 {id, name, type}，未命中返回 None。
    词典优先后无需区分 learned（学习映射不再需要内联防误判），外键保证目标必存在。
    带出 type：目标是「退款/不计收支」组时整行转不计收支。"""
    if not source:
        return None
    r = conn.execute(
        "SELECT m.category_id, c.name AS name, c.type AS type FROM category_mappings m "
        "JOIN categories c ON m.category_id=c.id "
        "WHERE m.ledger_id=? AND (m.flow_type=? OR c.type='transfer') AND m.source=?",
        (ledger_id, flow_type, source)).fetchone()
    return {"id": r["category_id"], "name": r["name"], "type": r["type"]} if r else None

_BANK_RE = re.compile(r"银行|储蓄卡|信用卡|借记卡|贷记卡")

# 平台钱包关键词（含花呗/零钱通等信用与理财钱包，按用户规则归平台账户）
_WALLET_KEYWORDS = {
    "alipay": ("余额宝", "账户余额", "花呗", "支付宝"),
    "wechat": ("零钱", "微信"),
    "jd": ("余额", "白条", "京东"),
}
_PLATFORM_ACC = {"alipay": "支付宝", "wechat": "微信支付", "jd": "京东"}

def _ensure_account(conn: sqlite3.Connection, ledger_id: int, name: str, atype: str = "bank") -> int:
    """账户不存在则自动创建（导入解析期即时生效，独立提交；撤销导入不删，后续自动复用）。"""
    row = conn.execute("SELECT id FROM accounts WHERE ledger_id=? AND name=?",
                       (ledger_id, name)).fetchone()
    if row:
        return row["id"]
    with conn:
        sort = conn.execute("SELECT COALESCE(MAX(sort),0)+1 FROM accounts WHERE ledger_id=?",
                            (ledger_id,)).fetchone()[0]
        cur = conn.execute("INSERT INTO accounts(ledger_id, name, type, sort) VALUES(?,?,?,?)",
                           (ledger_id, name, atype, sort))
    return cur.lastrowid

def _match_account(conn: sqlite3.Connection, ledger_id: int, platform: str, pay: str) -> int | None:
    """收/付款方式 → 账户 id（未匹配返回 None）：
    ① 银行卡段优先（& 组合取左起第一个银行卡段），按原文自动建户（如 招商银行储蓄卡(7529)）；
    ② 平台钱包/空值/" 归平台账户（花呗→支付宝、零钱通→微信支付），平台账户缺失自动补建（京东）；
    ③ 其余（本应用导出等）按现有账户名子串匹配。"""
    parts = [p.strip() for p in (pay or "").split("&") if p.strip() and p.strip() != "/"]
    bank = next((p for p in parts if _BANK_RE.search(p)), None)
    if bank:
        return _ensure_account(conn, ledger_id, bank, "bank")
    if platform in _PLATFORM_ACC and (
            not parts or any(any(k in p for k in _WALLET_KEYWORDS[platform]) for p in parts)):
        return _ensure_account(conn, ledger_id, _PLATFORM_ACC[platform], "e-wallet")
    accs = {r["name"]: r["id"] for r in conn.execute(
        "SELECT id, name FROM accounts WHERE ledger_id=?", (ledger_id,))}
    for p in parts[:1]:
        for name, aid in accs.items():
            if name in p or p in name:
                return aid
    return None

def _account_name(conn: sqlite3.Connection, acc_id: int | None) -> str:
    """根据 account_id 反查账户名；id 空则返回空串。"""
    if not acc_id:
        return ""
    r = conn.execute("SELECT name FROM accounts WHERE id=?", (acc_id,)).fetchone()
    return r["name"] if r else ""

# ---------- 预览 ----------

def _purge_bad_learned_mappings(conn: sqlite3.Connection) -> int:
    """修复历史 AI 误判写入的脏学习映射（v6 单行模型）：
    ① 与内置默认同源但分类不同的学习行 → 删除后按默认补种；
    ② 与默认完全同源同分类的学习行 → 还原 learned=0（它就是默认）；
    ③ 补种缺失的默认映射。"""
    removed = 0
    with conn:
        for lid in [r["id"] for r in conn.execute("SELECT id FROM ledgers")]:
            for flow, source, target in dbm.DEFAULT_MAPPINGS:
                cid = dbm.top_cat_id(conn, lid, target, flow)
                if not cid:
                    continue
                row = conn.execute(
                    "SELECT id, category_id, learned FROM category_mappings "
                    "WHERE ledger_id=? AND flow_type=? AND source=?", (lid, flow, source)).fetchone()
                if row and row["learned"]:
                    if row["category_id"] != cid:
                        conn.execute("DELETE FROM category_mappings WHERE id=?", (row["id"],))
                        removed += 1
                        row = None
                    else:
                        conn.execute("UPDATE category_mappings SET learned=0 WHERE id=?", (row["id"],))
                        continue
                if not row:
                    conn.execute(
                        "INSERT OR IGNORE INTO category_mappings(ledger_id, flow_type, source, category_id) "
                        "VALUES(?,?,?,?)", (lid, flow, source, cid))
    return removed

def auto_fix_transactions(conn: sqlite3.Connection, ledger_id: int) -> int:
    """修复历史误分类流水：行文本（名称+商户+备注）命中关键词词典、
    但当前一级分类与词典矛盾（如停车费被误分为餐饮美食）→ 纠正到词典分类。
    仅按确定性关键词修，不碰转账/退款行。"""
    rows = conn.execute(
        "SELECT t.id, t.type, t.amount, t.name, t.merchant, t.note, p.name AS top_name "
        "FROM transactions t JOIN categories c ON t.category_id=c.id "
        "LEFT JOIN categories p ON c.parent_id=p.id "
        "WHERE t.ledger_id=? AND t.type IN ('expense','income') AND t.refund_for IS NULL",
        (ledger_id,)).fetchall()
    fixes = {}
    for r in rows:
        kw = _top_by_keyword(f"{r['name'] or ''} {r['merchant'] or ''} {r['note'] or ''}", r["amount"])
        if kw and kw != r["top_name"]:
            fixes.setdefault((kw, r["type"]), []).append(r["id"])
    n = 0
    with conn:
        for (top_name, ftype), ids in fixes.items():
            cat = conn.execute(
                "SELECT id FROM categories "
                "WHERE ledger_id=? AND parent_id IS NULL AND name=? AND type=?",
                (ledger_id, top_name, ftype)).fetchone()
            if not cat:
                continue
            q = ",".join("?" * len(ids))
            cur = conn.execute(
                f"UPDATE transactions SET category_id=? WHERE id IN ({q})", (cat["id"], *ids))
            n += cur.rowcount
    return n

def build_previews(conn: sqlite3.Connection, ledger_id: int, flows: list[RawFlow],
                   learn_userfix: dict | None = None) -> list[dict]:
    """RawFlow → FlowPreview（含分类解析/去重/退款匹配预填，7.5/7.8）。"""
    _purge_bad_learned_mappings(conn)
    auto_fix_transactions(conn, ledger_id)
    # 批量去重（7.8 ②：一次查库建字典）
    hashes = {_row_hash(f["hash_fields"]) for f in flows if "hash_fields" in f}
    existing = {}
    if hashes:
        q = ",".join("?" * len(hashes))
        for r in conn.execute(
                f"SELECT id, source, row_hash FROM transactions WHERE ledger_id=? AND row_hash IN ({q})",
                (ledger_id, *hashes)):
            existing[(r["source"], r["row_hash"])] = r["id"]

    cats = list(conn.execute(
        "SELECT c.id, c.name, c.type, c.parent_id, p.name AS pname FROM categories c "
        "LEFT JOIN categories p ON c.parent_id=p.id WHERE c.ledger_id=?", (ledger_id,)))
    cat_top = {(c["name"], c["type"]): c for c in cats if not c["pname"]}
    _tcat = {"退款": dbm.top_cat_id(conn, ledger_id, "退款", "transfer"),
             "不计收支": dbm.top_cat_id(conn, ledger_id, "不计收支", "transfer")}

    def _top(name: str, ftype: str) -> dict | None:
        """按 (名字, 类型) 精确查一级分类——同名收支分类（人情往来）互不覆盖。"""
        return cat_top.get((name, ftype)) or cat_top.get((name, ""))

    def _fallback(ftype: str) -> dict | None:
        """兜底一级分类：支出→「其他」，收入→「其他收入」。"""
        return _top("其他" if ftype == "expense" else "其他收入", ftype)

    previews = []
    for f in flows:
        if "skip" in f or "hash_fields" not in f:
            previews.append({"skip": True, "reason": f.get("skip", "无效行"),
                             "name": f.get("name", ""), "amount": 0})
            continue
        h = _row_hash(f["hash_fields"])
        p = {"type": f["type"], "amount": f["amount"], "date": f["date"],
             "name": f["name"], "merchant": f.get("merchant", ""), "note": f.get("note", ""),
             "platform": f["platform"], "source": f["platform"],
             "order_no": f.get("order_no", ""), "row_hash": h,
             "category_id": None, "category_name": "", "parent_category": "",
             "classified_via": "raw", "ai_confidence": None,
             "need_confirm": f.get("need_confirm", False), "refund_for": None, "refund_status": "none",
             "refund_confidence": None,  # high/low/unmatched，仅 is_refund=True 才有意义
             "duplicate": False, "duplicate_of": None,
             "to_account_id": None, "tags": {}, "ai_meta": {}}
        dup = existing.get((f["platform"], h))
        if dup:
            p["duplicate"], p["duplicate_of"] = True, dup
            previews.append(p)
            continue

        if f.get("origin_source"):
            p["source"] = f["origin_source"]

        acc_id = _match_account(conn, ledger_id, f.get("platform", ""), f.get("pay_account", ""))
        p["account_id"] = acc_id
        p["account_name"] = _account_name(conn, acc_id) if acc_id else ""

        if f["type"] == "transfer":
            # 导出回环：收款账户列按名解析（先精确后建户），转账对方不丢
            if f.get("to_account_name"):
                row = conn.execute("SELECT id FROM accounts WHERE ledger_id=? AND name=?",
                                   (ledger_id, f["to_account_name"])).fetchone()
                p["to_account_id"] = row["id"] if row else _ensure_account(
                    conn, ledger_id, f["to_account_name"], "bank")
            sub = dbm.TRANSFER_SUBCLASS.get(f.get("source_category", ""), "账户转账")
            if f.get("is_refund") or f.get("export_refund_idx") is not None:
                sub = "退款"
            p["tags"]["transfer_sub"] = sub
            p["category_name"] = sub
            p["classified_via"] = "transfer"
            p["category_id"] = _tcat["退款" if sub == "退款" else "不计收支"]
            if (not f.get("is_refund") and f.get("export_refund_idx") is None
                    and f.get("source_category") not in dbm.TRANSFER_SUBCLASS
                    and f.get("platform") != "export"):   # 导出回环的转账源端已确认，不重标
                p["need_confirm"] = True
            # 导出回环：退款对按流水ID批内直连（负序号走 apply_import 回填），高置信免确认
            if f.get("export_refund_idx") is not None:
                tgt = flows[-f["export_refund_idx"] - 1]
                p["refund_for"] = f["export_refund_idx"]
                p["refund_confidence"] = "high"
                p["refund_status"] = "full" if tgt["amount"] == f["amount"] else "partial"
                p["need_confirm"] = False
        else:
            # 来源就叫「人情往来」（导出格式）时尊重原分类，跳过词典的人情重切
            text = f"{f.get('name','')} {f.get('merchant','')} {f.get('note','')}"
            raw_cat = (f.get("source_category") or "").strip()
            src_is_renqing = raw_cat == "人情往来" and bool(_top(raw_cat, f["type"]))
            # （组合支付备注里的「&红包」字样不该把已确认的电费改判成红包礼金）
            export_direct = (f.get("platform") == "export" and raw_cat
                             and _top(raw_cat, f["type"]))
            kw = None if (src_is_renqing or export_direct) else _top_by_keyword(text, f["amount"])
            kw_top = _top(kw, f["type"]) if kw else None
            if kw_top:
                top = kw_top
                p["classified_via"] = "keyword"
                if top["name"] == "红包礼金":
                    p["need_confirm"] = dbm.renqing_need_confirm(text)
            else:
                # ② 映射表（v6 外键直取分类行，结构上无悬空）
                target = _resolve_by_mapping(conn, ledger_id, f["type"], raw_cat)
                direct = export_direct or (_top(raw_cat, f["type"]) if raw_cat else None)
                if target and target["type"] == "transfer":
                    # 用户显式指定 → 高置信不进待确认；top 设好后走公共尾行（人情复核天然跳过）
                    sub = "退款" if target["name"] == "退款" else dbm.TRANSFER_SUBCLASS.get(raw_cat, "账户转账")
                    p["type"] = "transfer"
                    p["tags"]["transfer_sub"] = sub
                    p["need_confirm"] = False
                    p["classified_via"] = "mapping"
                    top = {"id": target["id"], "name": sub}
                elif export_direct:
                    top, p["classified_via"] = export_direct, "direct-name"
                elif target:
                    top = target
                    p["classified_via"] = "mapping"
                # ②' 同名直连：来源分类名与账本现有分类同名（本应用导出的分类名即同款）→ 直接归类
                elif direct:
                    top = direct
                    p["classified_via"] = "direct-name"
                else:
                    # ③ AI 兜底（由上层 ai 模块在预览后补充），④ 硬兜底「其他/其他收入」
                    top = _fallback(f["type"])
                    p["classified_via"] = "raw"
                    p["need_confirm"] = True
            p["category_id"] = top["id"] if top else None
            p["parent_category"] = top["name"] if top else (
                raw_cat or ("其他" if f["type"] == "expense" else "其他收入"))
            # 记录来源分类名：用户在预览里纠正分类后回写学习映射的匹配键
            if raw_cat:
                p["map_source"] = raw_cat
            p["category_name"] = p["parent_category"]
            # 人情类复核：词典分支已按金额切分过，这里只复核映射/同名直连的结果；
            if (p["classified_via"] != "keyword" and raw_cat != "人情往来"
                    and f.get("platform") != "export"
                    and p["parent_category"] in ("人情往来", "红包礼金")):
                alt = _top(_renqing_split(f"{raw_cat} {text}", f["amount"]), f["type"])
                if alt:
                    p["category_id"] = alt["id"]
                    p["parent_category"] = p["category_name"] = alt["name"]
                # 切分结果是「人情往来」直接入库，是「红包礼金」按统一口径决定是否待确认
                p["need_confirm"] = (p["parent_category"] == "红包礼金"
                                     and dbm.renqing_need_confirm(text))
        if f.get("extra_categories"):
            p["tags"]["extra_categories"] = f["extra_categories"]
        if f.get("attribution"):
            p["tags"]["attribution"] = f["attribution"]
        previews.append(p)

    # 退款匹配预填（7.7）
    _match_refunds(conn, ledger_id, previews, flows)
    # 精确去重）/批内重复 → dup_hint 由前端标「疑似重复」；精确 duplicate 仍自动跳过，互不影响
    dbm.dup_hints(conn, ledger_id, [p for p in previews if not p.get("skip")
                                    and not p.get("duplicate") and "amount" in p])
    return previews

def _common_prefix_ratio(a: str, b: str) -> float:
    """最长公共前缀 / 较短串长（京东改写标题的模糊度量，7.7 ② v1.2）。"""
    a, b = a[:60], b[:60]
    n = min(len(a), len(b))
    i = 0
    while i < n and a[i] == b[i]:
        i += 1
    return i / max(min(len(a), len(b)), 1)

def _match_refunds(conn: sqlite3.Connection, ledger_id: int, previews: list[dict], flows: list[RawFlow]) -> None:
    """分层匹配：①订单号 ②前缀/模糊 ②'微信 ③启发式。低置信 → need_confirm。"""
    expenses = [{"id": e["id"], "name": e["name"] or "", "merchant": e["merchant"] or "",
                 "amount": e["amount"], "order_no": e["order_no"] or "",
                 "date": e["date"], "source": e["source"]} for e in conn.execute(
        "SELECT id, name, merchant, amount, order_no, date, source FROM transactions "
        "WHERE ledger_id=? AND type='expense'", (ledger_id,))]
    # 本批次支出预览也作为候选（退款与原消费常在同一个账单文件里）；负数 id = -(预览序号+1)
    for i, e in enumerate(previews):
        if e.get("type") == "expense" and not e.get("skip") and not e.get("duplicate"):
            expenses.append({"id": -(i + 1), "name": e.get("name") or "", "merchant": e.get("merchant") or "",
                             "amount": e["amount"], "order_no": e.get("order_no") or "",
                             "date": e["date"], "source": e.get("source") or ""})
    for p, f in zip(previews, flows):
        if "skip" in p or not f.get("is_refund"):
            continue
        # 退款之后产生的同单号/同名/同额订单不进候选，与手动关联校验一致；逐层落空则
        # need_confirm 待人工核对
        pday = (p["date"] or "")[:10]
        pool = [e for e in expenses if (e["date"] or "")[:10] <= pday]
        # 京东退款行「交易说明」可能带"退款-"/"【退款】"/"退款·"前缀，也可能与原消费完全同名
        strip_prefixes = ("退款-", "退款·", "【退款】", "退款:")
        base_candidates = [p["name"]]
        for pref in strip_prefixes:
            if p["name"].startswith(pref) and len(p["name"]) > len(pref):
                base_candidates.append(p["name"][len(pref):])
        base_candidates = list(dict.fromkeys(c for c in base_candidates if c))  # 去重保序
        base_candidates = [c for c in base_candidates if not _is_placeholder_name(c)]
        cand = None
        # ① 订单号精确
        for e in pool:
            if p["order_no"] and e["order_no"] and p["order_no"] == e["order_no"]:
                cand = (e, False)
                break
        if cand is None:
            best: tuple | None = None   # (expense, low_conf, ratio)
            for e in pool:
                if e["amount"] < p["amount"]:
                    continue
                # 金额相等 + 任一 base 是 e.name 的前缀 → 高置信
                if e["amount"] == p["amount"]:
                    for base in base_candidates:
                        if e["name"].startswith(base):
                            best = (e, False, 1.0)
                            break
                    if best:
                        break
                # 模糊：最长公共前缀比 ≥ 0.5（任一 base 与 e.name）
                r = 0.0
                for base in base_candidates:
                    r = max(r, _common_prefix_ratio(base, e["name"]))
                if r >= 0.5 and (best is None or r > best[2]):
                    best = (e, True, r)
            if best:
                cand = (best[0], best[1])
        # ②' 微信：商户+金额 / 红包转账（类型+金额+最近）
        if cand is None and f["platform"] == "wechat":
            wt = f.get("wechat_type", "")
            rows = sorted([e for e in pool
                           if e["amount"] == p["amount"] and e["source"] not in ("manual", "app")],
                          key=lambda e: e["date"], reverse=True)
            if wt in ("微信红包", "转账"):
                if rows:
                    cand = (rows[0], True)  # 无商户：类型+金额+最近 → 低置信
            else:
                for e in rows:
                    if e["merchant"] and e["merchant"] == p["merchant"]:
                        cand = (e, False)
                        break
        # ③ 启发式：金额+日期在原消费后 0-90 天
        if cand is None:
            cands = [e for e in pool if e["amount"] >= p["amount"] and _within_days(e["date"], p["date"], 90)]
            if len(cands) == 1:
                cand = (cands[0], True)
        if cand:
            e, low_conf = cand
            p["refund_for"], p["refund_status"] = e["id"], ("partial" if e["amount"] > p["amount"] else "full")
            p["refund_confidence"] = "low" if low_conf else "high"
            if low_conf:
                p["need_confirm"] = True
        else:
            # 未匹配到原消费的退款（如被退回的红包）：进待确认队列，由用户核对
            p["need_confirm"] = True
            p["refund_confidence"] = "unmatched"

def _within_days(d1: str, d2: str, days: int) -> bool:
    try:
        a = datetime.strptime(d1[:19], "%Y-%m-%d %H:%M:%S")
        b = datetime.strptime(d2[:19], "%Y-%m-%d %H:%M:%S")
        return 0 <= (b - a).days <= days
    except ValueError:
        return False

# ---------- 入库 ----------

def apply_import(conn: sqlite3.Connection, ledger_id: int, previews: list[dict],
                 filename: str, source: str, learn: bool = True) -> int:
    """确认导入（7.8 ⑤~⑦）：跳过 duplicate/skip，建批次，写新行，learn 回写。"""
    new = [p for p in previews if not p.get("duplicate") and not p.get("skip")]
    dates = [p["date"] for p in new]
    with conn:
        cur = conn.execute(
            "INSERT INTO import_batches(ledger_id, source, filename, count, imported_at, date_from, date_to) "
            "VALUES(?,?,?,?,datetime('now'),?,?)",
            (ledger_id, source, filename, len(new), min(dates) if dates else "", max(dates) if dates else ""))
        batch_id = cur.lastrowid
        inserted_ids = []
        pending_refund: list[tuple[int, int]] = []  # (本行序号, 目标预览序号)：负数 id 待整批插完后回填
        for p in new:
            # 本批次内部退款关联：负数 id 指向预览序号，先记下待回填——退款行常排在原消费之前
            # （微信账单按时间倒序，退款晚于原消费故更靠前），插入时目标行往往还没入库，
            rf = p.get("refund_for")
            if rf is not None and rf < 0:
                pending_refund.append((len(inserted_ids), -rf - 1))
                rf = None
            # 来源分类名存入 ai_meta：待确认队列改分类时用它做学习键（旧版用 note，导入行常为空）
            if p.get("map_source"):
                meta = p.get("ai_meta") or {}
                meta["map_source"] = p["map_source"]
                p["ai_meta"] = meta
            # 之前复用 cur（指向 import_batches），本批次内退款关联（京东拆分必走）失效
            cur = conn.execute(
                """INSERT INTO transactions(ledger_id, type, amount, batch_id, category_id,
                   account_id, to_account_id, date, note, name, merchant, tags, receipt_path,
                   source, ai_meta, order_no, refund_for, refund_status, refund_confidence,
                   need_confirm, row_hash)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (ledger_id, p["type"], p["amount"], batch_id, p.get("category_id"),
                 p.get("account_id"), p.get("to_account_id"), p["date"], p.get("note", ""),
                 p.get("name", ""), p.get("merchant", ""), dbm._j(p.get("tags")), None,
                 p.get("source"), dbm._j(p.get("ai_meta")),
                 p.get("order_no", ""), rf, p.get("refund_status", "none"),
                 p.get("refund_confidence"),  # NULL=非退款 / high / low / unmatched
                 1 if p.get("need_confirm") else 0, p["row_hash"]))
            inserted_ids.append(cur.lastrowid)
            # learn：只学习用户在预览里手动改过的分类（user_changed 由前端 impSetCat 标记）。
            # 自动结果(mapping/direct-name/keyword/AI)一律不回写——AI 兜底仅单次生效于当批流水；
            # 用户手动纠正过的行(含纠正 AI 行)是用户意图，允许回写。
            # v6：目标直接用 category_id（旧版用 parent_category 文本，前端改分类不更新该字段
            # 会学到旧分类名）；且同源唯一一行，统一走 dbm.upsert_mapping。
            if learn and p.get("user_changed") and p.get("map_source") and p.get("category_id"):
                dbm.upsert_mapping(conn, ledger_id, p["type"], p["map_source"],
                                   p["category_id"], batch_id=batch_id)
        # 回填批内退款关联（顺序无关）：负数预览序号 → 真实行 id。
        # 目标不在本批（序号越界）则保持 NULL，交给下面的反向补关联（7.7 匹配时机②）
        for row_pos, tgt_pos in pending_refund:
            if tgt_pos < len(inserted_ids):
                conn.execute("UPDATE transactions SET refund_for=? WHERE id=?",
                             (inserted_ids[tgt_pos], inserted_ids[row_pos]))
        # 反向补关联：本批新 expense 可能是此前未匹配退款行的原消费（7.7 匹配时机②）
        _rescan_unmatched_refunds(conn, ledger_id)
    return batch_id

def _rescan_unmatched_refunds(conn: sqlite3.Connection, ledger_id: int) -> None:
    for r in conn.execute(
            "SELECT id, name, merchant, amount, order_no, date FROM transactions "
            "WHERE ledger_id=? AND type='transfer' AND refund_for IS NULL "
            "AND tags LIKE '%\"transfer_sub\": \"退款\"%'", (ledger_id,)).fetchall():
        base = r["name"][3:] if r["name"].startswith("退款-") else r["name"]
        # 19 位日期也放行；占位名分支原 date<=? 裸串比较会把同日带时间的原消费漏掉，一并修正）
        rday = (r["date"] or "")[:10]
        hit = None
        if r["order_no"]:
            hit = conn.execute(
                "SELECT id, amount FROM transactions WHERE ledger_id=? AND type='expense' AND order_no=? "
                "AND substr(date,1,10)<=? LIMIT 1",
                (ledger_id, r["order_no"], rday)).fetchone()
        if not hit and base:
            if _is_placeholder_name(base):
                # 微信红包/转账行「商品」列写作 "/"，名称不含信息，按名匹配会把退款静默关联
                # 到任意一条同为 "/" 的原消费（且不受金额约束）→ 改用 ②' 同口径：
                hit = conn.execute(
                    "SELECT id, amount FROM transactions WHERE ledger_id=? AND type='expense' "
                    "AND amount=? AND substr(date,1,10)<=? ORDER BY date DESC LIMIT 1",
                    (ledger_id, r["amount"], rday)).fetchone()
            else:
                hit = conn.execute(
                    "SELECT id, amount FROM transactions WHERE ledger_id=? AND type='expense' "
                    "AND (name=? OR name LIKE ?) AND substr(date,1,10)<=? LIMIT 1",
                    (ledger_id, base, base + "%", rday)).fetchone()
        if hit:
            status = "partial" if hit["amount"] > r["amount"] else "full"
            conn.execute("UPDATE transactions SET refund_for=?, refund_status=? WHERE id=?",
                         (hit["id"], status, r["id"]))

def undo_batch(conn: sqlite3.Connection, batch_id: int) -> int:
    """整批撤销（7.8 ⑦）：删行 + 级联清理退款关联/learned 映射 + 删批次。"""
    with conn:
        n = conn.execute("SELECT COUNT(*) FROM transactions WHERE batch_id=?", (batch_id,)).fetchone()[0]
        ids = [r[0] for r in conn.execute("SELECT id FROM transactions WHERE batch_id=?", (batch_id,))]
        if ids:
            q = ",".join("?" * len(ids))
            conn.execute(f"UPDATE transactions SET refund_for=NULL, refund_status='none', need_confirm=1 "
                         f"WHERE refund_for IN ({q})", ids)
            conn.execute("DELETE FROM transactions WHERE batch_id=?", (batch_id,))
        conn.execute("DELETE FROM category_mappings WHERE learned=1 AND batch_id=?", (batch_id,))
        conn.execute("DELETE FROM import_batches WHERE id=?", (batch_id,))
    return n

